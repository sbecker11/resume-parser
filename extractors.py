#!/usr/bin/env python3
"""
Extract raw text from resume documents (DOCX, PDF).
"""

import re
from pathlib import Path


def extract_text(path: str | Path) -> str:
    """
    Extract text from a resume file. Supports .docx and .pdf.
    Returns plain text with section structure preserved.
    """
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".docx":
        return _extract_docx(path)
    elif suffix == ".pdf":
        return _extract_pdf(path)
    else:
        raise ValueError(f"Unsupported format: {suffix}. Use .docx or .pdf")


def _extract_docx(path: Path) -> str:
    """Extract text from a DOCX file using python-docx."""
    from docx import Document

    doc = Document(path)
    parts = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n\n".join(parts)


def _extract_pdf(path: Path) -> str:
    """Extract text from a PDF file using pdfplumber."""
    import pdfplumber

    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text.strip())

    return "\n\n".join(parts)
